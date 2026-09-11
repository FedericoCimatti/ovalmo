# -*- coding: utf-8 -*-
"""Il foglio Excel, con le formule vive e il grafico dell'andamento.

E' il build.py di sempre: stesso codice, stessi colori, stesse formule. L'unica
differenza e' che il corpo sta dentro una funzione, cosi' il job automatico puo'
chiamarlo senza lanciare un altro programma. Sotto c'e' anche `verifica`, che
riapre il file e controlla che non ci siano formule rotte prima di committarlo.

Le righe vengono create solo per le giornate presenti in `calendario`:
aggiungendo una giornata al JSON, il file cresce da solo.
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.formatting.rule import CellIsRule, DataBarRule, ColorScaleRule

# le regole di gioco stanno in un posto solo: qui si leggono, non si riscrivono
from .punteggio import SOLITARIO_DA, SOLITARIO_ESATTO as SOL_ESATTO, SOLITARIO_SEGNO as SOL_SEGNO


def genera(DATA, OUT):
    """Scrive il file Excel. DATA e' la stagione unita ai pronostici."""
    PLAYERS = DATA["players"]
    MPG = DATA["partite_per_giornata"]
    CAL = {int(k): [tuple(x) for x in v] for k, v in DATA["calendario"].items()}
    RIS, PRO = DATA.get("risultati", {}), DATA.get("pronostici", {})
    GIORNATE = sorted(CAL)
    R0 = 3
    RN = R0 + len(GIORNATE) * MPG - 1
    NP = len(PLAYERS)

    F = "Arial"
    NAVY, NAVY2, GOLD, GOLD_D = "0F2544", "1B3A63", "C8A24A", "8A6D1F"
    INK, MUTE, LINE = "1A1A1A", "8A94A6", "D8DEE8"
    PANEL, PANEL2, WHITE = "F4F6FA", "EAEEF5", "FFFFFF"
    INPUT_BG, INPUT_FG = "FFF8E1", "1B4DB1"
    GREEN, AMBER = "D9F0DC", "FCEFC7"
    PCOL = ["DCE9F7", "DEF0E4", "FBE3D8", "E7E3F6", "FDF2D0"]
    PDOT = ["2E6FB7", "3E8E5A", "C2622F", "6A56A8", "B08A1E"]
    MEDAL = [("FFF3CF", "8A6D1F", "1o POSTO"), ("EDF0F4", "5A6472", "2o POSTO"), ("F7E7DA", "8A5A34", "3o POSTO")]

    hair = Side(style="thin", color=LINE)
    BOX = Border(left=hair, right=hair, top=hair, bottom=hair)
    BOT = Border(bottom=Side(style="thin", color=LINE))

    def s(c, *, b=False, sz=10, fg=INK, bg=None, al="left", wrap=False, fmt=None, it=False, box=None, h=None):
        c.font = Font(name=F, bold=b, size=sz, color=fg, italic=it)
        if bg: c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
        if fmt: c.number_format = fmt
        if box is not None: c.border = box
        return c

    def inp(c, sz=10):
        return s(c, b=True, sz=sz, fg=INPUT_FG, bg=INPUT_BG, al="center", box=BOX)

    def band(ws, rng, a1, a2, *, bg, fg, txt, sz, al="left", h=None):
        ws.merge_cells(rng)
        s(ws[a1], b=True, sz=sz, fg=fg, bg=bg, al=al); ws[a1] = txt
        for col in range(ws[a1].column, ws[a2].column + 1):
            ws.cell(row=ws[a1].row, column=col).fill = PatternFill("solid", fgColor=bg)
        if h: ws.row_dimensions[ws[a1].row].height = h

    def widths(ws, spec):
        for col, w in spec.items(): ws.column_dimensions[col].width = w

    wb = Workbook()

    # =========================================================== CLASSIFICA
    cl = wb.active; cl.title = "Classifica"
    cl.sheet_view.showGridLines = False; cl.sheet_properties.tabColor = GOLD
    GEN, MR0 = 12, 0
    MRN = 0

    widths(cl, {"A": 6, "B": 19, "C": 11, "D": 11, "E": 11, "F": 12, "G": 12, "H": 11,
                "I": 23, "J": 2, "K": 17, "L": 13, "M": 13, "N": 13, "O": 13, "P": 13})

    band(cl, "A1:I2", "A1", "I2", bg=NAVY, fg=WHITE, txt="TROFEO OVALMO", sz=24, h=34)
    cl.merge_cells("K1:P2")
    s(cl["K1"], b=True, sz=11, fg=WHITE, bg=NAVY, al="right"); cl["K1"] = "Serie A 2026/27"
    for col in range(11, 17): cl.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NAVY)
    for col in range(1, 17): cl.cell(row=2, column=col).fill = PatternFill("solid", fgColor=NAVY)
    cl.row_dimensions[2].height = 6
    band(cl, "A3:P3", "A3", "P3", bg=GOLD, fg=NAVY,
         txt="   Classifica aggiornata in automatico - non serve toccare nulla in questo foglio", sz=9, h=17)

    # --- podio ---
    spans = [("B", "C"), ("E", "F"), ("H", "I")]
    for i, ((c1, c2), (bg, fg, lab)) in enumerate(zip(spans, MEDAL)):
        cl.merge_cells(f"{c1}5:{c2}5")
        s(cl[f"{c1}5"], b=True, sz=9, fg=fg, bg=bg, al="center", box=BOX); cl[f"{c1}5"] = lab
        cl[f"{c2}5"].fill = PatternFill("solid", fgColor=bg); cl[f"{c2}5"].border = BOX
        cl.merge_cells(f"{c1}6:{c2}6")
        cl[f"{c1}6"] = (f'=IFERROR(INDEX($B${GEN}:$B${GEN+NP-1},MATCH({i+1},$A${GEN}:$A${GEN+NP-1},0))'
                        f'&IF(COUNTIF($A${GEN}:$A${GEN+NP-1},{i+1})>1," (a pari)",""),"-")')
        s(cl[f"{c1}6"], b=True, sz=15, fg=NAVY, bg=bg, al="center", box=BOX)
        cl[f"{c2}6"].fill = PatternFill("solid", fgColor=bg); cl[f"{c2}6"].border = BOX
        cl.merge_cells(f"{c1}7:{c2}7")
        cl[f"{c1}7"] = (f'=IFERROR(INDEX($C${GEN}:$C${GEN+NP-1},MATCH({i+1},$A${GEN}:$A${GEN+NP-1},0))'
                        f'&" punti","")')
        s(cl[f"{c1}7"], sz=9, fg=fg, bg=bg, al="center", box=BOX)
        cl[f"{c2}7"].fill = PatternFill("solid", fgColor=bg); cl[f"{c2}7"].border = BOX
        cl.row_dimensions[6].height = 24

    # --- riquadri stato ---
    stat_defs = [("Giornate completate", None), ("Giornate iniziate", None),
                 ("Ultima completata", None), ("Prossima da giocare", None)]
    for i, (lab, _) in enumerate(stat_defs):
        r = 5 + i
        s(cl.cell(row=r, column=11, value=lab), sz=9, fg=MUTE, bg=PANEL, box=BOX)
        s(cl.cell(row=r, column=12), b=True, sz=12, fg=NAVY, bg=PANEL, al="center", box=BOX)

    s(cl["A10"], b=True, sz=13, fg=NAVY); cl["A10"] = "CLASSIFICA GENERALE"
    heads = ["Pos", "Giocatore", "Punti", "Segni presi", "Risultati esatti",
             "Pronostici inviati", "Media / giornata", "Giornate vinte"]
    for j, h in enumerate(heads):
        s(cl.cell(row=11, column=1+j, value=h), b=True, sz=9, fg=WHITE, bg=NAVY2,
          al="center", wrap=True, box=BOX)
    cl.row_dimensions[11].height = 30

    MR0 = 33
    MRN = MR0 + len(GIORNATE) - 1
    MB = f"$B${MR0}:$B${MRN}"; MA = f"$A${MR0}:$A${MRN}"
    widths(cl, {"K": 16, "L": 11, "M": 11, "N": 16, "O": 9, "P": 6})

    for i, f in enumerate([
        f'=COUNTIF({MB},"Completata")',
        f'=COUNTIF({MB},"Completata")+COUNTIF({MB},"In corso")',
        f'=IF($L$5=0,"-",SUMPRODUCT(MAX(({MB}="Completata")*{MA})))',
        f'=IF(SUMPRODUCT(MIN(({MB}<>"Completata")*{MA}+({MB}="Completata")*999))>=999,"-",'
        f'SUMPRODUCT(MIN(({MB}<>"Completata")*{MA}+({MB}="Completata")*999)))']):
        cl.cell(row=5+i, column=12, value=f)

    # --- tabella generale ---
    for i in range(NP):
        r = GEN + i
        pc = CL(4+i)          # colonna del giocatore nel foglio Punti
        mc = CL(3+i)          # colonna del giocatore nella matrice giornate
        sc, hc = CL(5+i*3), CL(6+i*3)
        cl.cell(row=r, column=1, value=f'=SUMPRODUCT((($C${GEN}:$C${GEN+NP-1}*1000+$E${GEN}:$E${GEN+NP-1})>($C{r}*1000+$E{r}))*1)+1')
        cl.cell(row=r, column=2, value=f"=Config!$A${6+i}")
        cl.cell(row=r, column=3, value=f"=SUM(Punti!${pc}${R0}:${pc}${RN})")
        colonna = f"Punti!${pc}${R0}:${pc}${RN}"
        # i segni giusti sono quattro casi: 1 e 3 di sempre, 2 e 6 del solitario
        cl.cell(row=r, column=4, value="=" + "+".join(
            f"COUNTIF({colonna},Config!$D${d})" for d in (7, 6, 10, 9)))
        cl.cell(row=r, column=5, value=f"=COUNTIF({colonna},Config!$D$6)+COUNTIF({colonna},Config!$D$9)")
        cl.cell(row=r, column=6, value=f'=SUMPRODUCT(((Pronostici!${sc}${R0}:${sc}${RN}<>"")+(Pronostici!${hc}${R0}:${hc}${RN}<>""))*1>0)')
        cl.cell(row=r, column=7, value=f'=IF($L$6=0,0,$C{r}/$L$6)')
        cl.cell(row=r, column=8, value=f'=SUMPRODUCT((${mc}${MR0}:${mc}${MRN}=$H${MR0}:$H${MRN})*($H${MR0}:$H${MRN}>0)*({MB}="Completata"))')
        for col in range(1, 9):
            s(cl.cell(row=r, column=col), b=(col in (1, 2, 3)), sz=11 if col == 3 else 10,
              fg=NAVY if col == 3 else INK, al="left" if col == 2 else "center",
              bg=WHITE if i % 2 == 0 else PANEL, box=BOX, fmt="0.00" if col == 7 else None)
        cl.cell(row=r, column=2).fill = PatternFill("solid", fgColor=PCOL[i])
        cl.row_dimensions[r].height = 20
    cl.conditional_formatting.add(f"C{GEN}:C{GEN+NP-1}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color=GOLD, showValue=True))
    s(cl.cell(row=GEN+NP, column=1), sz=8, fg=MUTE, it=True)
    cl.cell(row=GEN+NP, column=1).value = "A parita' di punti passa avanti chi ha piu' risultati esatti. Le giornate vinte contano solo quelle completate."
    cl.merge_cells(start_row=GEN+NP, start_column=1, end_row=GEN+NP, end_column=8)

    # --- pannello ultima giornata ---
    s(cl["A18"], b=True, sz=12, fg=NAVY); cl["A18"] = "ULTIMA GIORNATA COMPLETATA"
    cl.merge_cells("C18:D18")
    cl["C18"] = '=IF($L$7="-","nessuna ancora","Giornata "&$L$7)'
    s(cl["C18"], b=True, sz=11, fg=GOLD_D, al="left")
    for j, h in enumerate(["Giocatore", "Punti", ""]):
        s(cl.cell(row=19, column=2+j, value=h), b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX)
    for i in range(NP):
        r = 20 + i
        mc = CL(3+i)
        cl.cell(row=r, column=2, value=f"=Config!$A${6+i}")
        cl.cell(row=r, column=3, value=f'=IF($L$7="-","-",IFERROR(INDEX(${mc}${MR0}:${mc}${MRN},MATCH($L$7,{MA},0)),"-"))')
        cl.cell(row=r, column=4, value=f'=IF($L$7="-","",IF(IFERROR(INDEX(${mc}${MR0}:${mc}${MRN},MATCH($L$7,{MA},0)),0)=IFERROR(INDEX($H${MR0}:$H${MRN},MATCH($L$7,{MA},0)),-1),"RE DELLA GIORNATA",""))')
        s(cl.cell(row=r, column=2), b=True, sz=10, bg=PCOL[i], box=BOX)
        s(cl.cell(row=r, column=3), b=True, sz=11, fg=NAVY, al="center", bg=WHITE, box=BOX)
        s(cl.cell(row=r, column=4), b=True, sz=8, fg=GOLD_D, al="center", bg=WHITE, box=BOX)
        cl.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        cl.cell(row=r, column=5).border = BOX

    # --- pannello prossima giornata ---
    s(cl["K18"], b=True, sz=12, fg=NAVY); cl["K18"] = "PROSSIMA GIORNATA"
    cl.merge_cells("K19:M19")
    cl["K19"] = '=IF($L$8="-","Stagione completata","Giornata "&$L$8&" - partite da pronosticare")'
    s(cl["K19"], b=True, sz=9, fg=GOLD_D)
    PA, PB, PE, PF = f"Partite!$A${R0}:$A${RN}", None, f"Partite!$E${R0}:$E${RN}", f"Partite!$F${R0}:$F${RN}"
    for i in range(MPG):
        r = 20 + i
        key = f'"G"&TEXT($L$8,"00")&"-"&TEXT({i+1},"00")'
        cl.cell(row=r, column=11, value=(
            f'=IFERROR(INDEX({PE},MATCH({key},{PA},0))&"   -   "&INDEX({PF},MATCH({key},{PA},0)),"-")'))
        cl.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
        for col in (11, 12, 13):
            s(cl.cell(row=r, column=col), sz=10, bg=WHITE if i % 2 == 0 else PANEL, box=BOX)
        cl.cell(row=r, column=11).alignment = Alignment(horizontal="center", vertical="center")
    s(cl.cell(row=19, column=14, value="Pronostici"), b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX)
    s(cl.cell(row=19, column=15, value="ok"), b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX)
    for i in range(NP):
        r = 20 + i
        sc, hc = CL(5+i*3), CL(6+i*3)
        cl.cell(row=r, column=14, value=f"=Config!$A${6+i}")
        cl.cell(row=r, column=15, value=(
            f'=IF($L$8="-","-",SUMPRODUCT((Pronostici!$B${R0}:$B${RN}=$L$8)*'
            f'(((Pronostici!${sc}${R0}:${sc}${RN}<>"")+(Pronostici!${hc}${R0}:${hc}${RN}<>""))>0))&"/{MPG}")'))
        s(cl.cell(row=r, column=14), b=True, sz=9, bg=PCOL[i], box=BOX)
        s(cl.cell(row=r, column=15), b=True, sz=9, fg=NAVY, al="center", bg=WHITE, box=BOX)
    s(cl.cell(row=25, column=14), sz=8, fg=MUTE, it=True)
    cl.cell(row=25, column=14).value = "chi e' sotto 10 non ha finito"
    cl.merge_cells("N25:O26")

    # --- matrice punti per giornata ---
    s(cl[f"A31"], b=True, sz=12, fg=NAVY); cl["A31"] = "PUNTI PER GIORNATA"
    for j, h in enumerate(["G.", "Stato"]):
        s(cl.cell(row=32, column=1+j, value=h), b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX)
    for i in range(NP):
        c = cl.cell(row=32, column=3+i, value=f"=Config!$A${6+i}")
        s(c, b=True, sz=9, fg=NAVY, bg=PCOL[i], al="center", box=BOX)
    s(cl.cell(row=32, column=8, value="Max"), b=True, sz=8, fg=MUTE, bg=PANEL2, al="center", box=BOX)
    s(cl.cell(row=32, column=9, value="Re della giornata"), b=True, sz=9, fg=NAVY, bg=GOLD, al="center", box=BOX)
    s(cl.cell(row=32, column=17, value="cumulati"), b=True, sz=8, fg=MUTE)
    for i in range(NP):
        s(cl.cell(row=32, column=18+i, value=f"=Config!$A${6+i}"), b=True, sz=8, fg=MUTE, bg=PANEL2, al="center")

    PB_ = f"Partite!$B${R0}:$B${RN}"; PK_ = f"Partite!$K${R0}:$K${RN}"; PE_ = f"Partite!$E${R0}:$E${RN}"
    for k, g in enumerate(GIORNATE):
        r = MR0 + k
        cl.cell(row=r, column=1, value=g)
        cl.cell(row=r, column=2, value=(
            f'=IF(COUNTIFS({PB_},$A{r},{PE_},"<>")=0,"Calendario mancante",'
            f'IF(COUNTIFS({PB_},$A{r},{PK_},"Giocata")=0,"Non giocata",'
            f'IF(COUNTIFS({PB_},$A{r},{PK_},"Giocata")=COUNTIFS({PB_},$A{r},{PE_},"<>"),"Completata","In corso")))'))
        for i in range(NP):
            pc = CL(4+i)
            cl.cell(row=r, column=3+i, value=f"=SUMIF(Punti!$B${R0}:$B${RN},$A{r},Punti!${pc}${R0}:${pc}${RN})")
        cl.cell(row=r, column=8, value=f"=MAX($C{r}:${CL(2+NP)}{r})")
        cl.cell(row=r, column=9, value=(
            f'=IF($H{r}<=0,"-",INDEX($C$32:${CL(2+NP)}$32,MATCH($H{r},$C{r}:${CL(2+NP)}{r},0))'
            f'&IF(COUNTIF($C{r}:${CL(2+NP)}{r},$H{r})>1," (a pari)","")'
            f'&IF($B{r}="Completata",""," - in corso"))'))
        for i in range(NP):
            mc, cc = CL(3+i), CL(18+i)
            cl.cell(row=r, column=18+i, value=f"={mc}{r}" if r == MR0 else f"={cc}{r-1}+{mc}{r}")
        bg = WHITE if k % 2 == 0 else PANEL
        for col in range(1, 10):
            s(cl.cell(row=r, column=col), b=(col in (1, 9)), sz=10,
              fg=GOLD_D if col == 9 else INK, al="left" if col in (2, 9) else "center", bg=bg, box=BOX)
        s(cl.cell(row=r, column=8), sz=8, fg=MUTE, al="center", bg=PANEL2, box=BOX)
        for i in range(NP):
            s(cl.cell(row=r, column=18+i), sz=8, fg=MUTE, al="center", bg=PANEL2)
    cl.conditional_formatting.add(f"C{MR0}:{CL(2+NP)}{MRN}",
        ColorScaleRule(start_type="num", start_value=0, start_color=WHITE,
                       end_type="max", end_color="BBDDBB"))
    cl.freeze_panes = "A4"

    ch = LineChart()
    ch.title = "Andamento punti cumulati"
    ch.style = 2
    ch.y_axis.title = "Punti"; ch.x_axis.title = "Giornata"
    ch.height, ch.width = 9.5, 27
    ch.add_data(Reference(cl, min_col=18, max_col=17+NP, min_row=32, max_row=MRN), titles_from_data=True)
    ch.set_categories(Reference(cl, min_col=1, min_row=MR0, max_row=MRN))
    for i, ser in enumerate(ch.series):
        ser.smooth = False
        ser.graphicalProperties.line.solidFill = PDOT[i % len(PDOT)]
        ser.graphicalProperties.line.width = 22000
        ser.marker = Marker(symbol="circle", size=6)
    CH_ROW = MRN + 3
    cl.add_chart(ch, f"A{CH_ROW}")

    # =========================================================== PARTITE
    pt = wb.create_sheet("Partite")
    pt.sheet_view.showGridLines = False; pt.sheet_properties.tabColor = NAVY2
    widths(pt, {"A": 9, "B": 8, "C": 12, "D": 11, "E": 16, "F": 16, "G": 10, "H": 10, "I": 8, "J": 11, "K": 20})
    band(pt, "A1:F1", "A1", "F1", bg=NAVY, fg=WHITE, txt="  CALENDARIO", sz=11, h=22)
    band(pt, "G1:H1", "G1", "H1", bg=GOLD, fg=NAVY, txt="RISULTATO REALE", sz=10)
    pt["G1"].alignment = Alignment(horizontal="center", vertical="center")
    band(pt, "I1:K1", "I1", "K1", bg=NAVY2, fg=WHITE, txt="  CALCOLATO IN AUTOMATICO", sz=10)
    for col, h in zip("ABCDEFGHIJK", ["ID", "G.", "Data", "Ora", "Casa", "Ospite", "Gol Casa",
                                      "Gol Ospite", "Segno", "Risultato", "Stato"]):
        s(pt[f"{col}2"], b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", wrap=True, box=BOX)
        pt[f"{col}2"] = h
    pt.row_dimensions[2].height = 26

    r = R0
    for k, g in enumerate(GIORNATE):
        bg = WHITE if k % 2 == 0 else PANEL
        for m in range(1, MPG+1):
            fx = CAL[g][m-1] if m-1 < len(CAL[g]) else None
            mid = f"G{g:02d}-{m:02d}"
            pt.cell(row=r, column=1, value=mid); pt.cell(row=r, column=2, value=g)
            if fx:
                for j in range(4): pt.cell(row=r, column=3+j, value=fx[j])
            if RIS.get(mid):
                pt.cell(row=r, column=7, value=RIS[mid][0]); pt.cell(row=r, column=8, value=RIS[mid][1])
            pt.cell(row=r, column=9,  value=f'=IF(OR($G{r}="",$H{r}=""),"",IF($G{r}>$H{r},"1",IF($G{r}=$H{r},"X","2")))')
            pt.cell(row=r, column=10, value=f'=IF(OR($G{r}="",$H{r}=""),"",$G{r}&"-"&$H{r})')
            pt.cell(row=r, column=11, value=f'=IF($E{r}="","Calendario da inserire",IF(OR($G{r}="",$H{r}=""),"Da giocare","Giocata"))')
            for col in range(1, 12):
                s(pt.cell(row=r, column=col), sz=10, al="left" if col in (5, 6, 11) else "center",
                  bg=bg, box=BOX, fg=MUTE if col in (1, 11) else INK, b=(col in (5, 6)))
            for col in (7, 8): inp(pt.cell(row=r, column=col))
            s(pt.cell(row=r, column=9), b=True, sz=10, fg=NAVY, al="center", bg=bg, box=BOX)
            r += 1
    pt.freeze_panes = "C3"; pt.auto_filter.ref = f"A2:K{RN}"
    pt.conditional_formatting.add(f"K{R0}:K{RN}",
        CellIsRule(operator="equal", formula=['"Giocata"'], fill=PatternFill("solid", fgColor=GREEN)))

    # =========================================================== PRONOSTICI
    pr = wb.create_sheet("Pronostici")
    pr.sheet_view.showGridLines = False; pr.sheet_properties.tabColor = GOLD_D
    widths(pr, {"A": 9, "B": 8, "C": 16, "D": 16})
    band(pr, "A1:D1", "A1", "D1", bg=NAVY, fg=WHITE, txt="  PARTITA", sz=11, h=22)
    for col, h in zip("ABCD", ["ID", "G.", "Casa", "Ospite"]):
        s(pr[f"{col}2"], b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX); pr[f"{col}2"] = h
    for i in range(NP):
        c0 = 5 + i*3
        pr.merge_cells(f"{CL(c0)}1:{CL(c0+2)}1")
        pr[f"{CL(c0)}1"] = f"=Config!$A${6+i}"
        s(pr[f"{CL(c0)}1"], b=True, sz=12, fg=NAVY, bg=PCOL[i], al="center", box=BOX)
        for j in range(3): pr.cell(row=1, column=c0+j).fill = PatternFill("solid", fgColor=PCOL[i])
        for j, h in enumerate(["Segno", "Gol C", "Gol O"]):
            s(pr.cell(row=2, column=c0+j, value=h), b=True, sz=8, fg=NAVY, bg=PCOL[i], al="center", box=BOX)
            pr.column_dimensions[CL(c0+j)].width = 7.5
    pr.row_dimensions[2].height = 18

    dv = DataValidation(type="list", formula1='"1,X,2"', allow_blank=True, showErrorMessage=True)
    dv.error = "Scrivi 1 (vince la casa), X (pareggio) oppure 2 (vince l'ospite)."
    dv.errorTitle = "Segno non valido"
    pr.add_data_validation(dv)

    r = R0
    for k, g in enumerate(GIORNATE):
        bg = WHITE if k % 2 == 0 else PANEL
        for m in range(1, MPG+1):
            mid = f"G{g:02d}-{m:02d}"
            guess = PRO.get(mid, {})
            pr.cell(row=r, column=1, value=f"=Partite!$A{r}")
            pr.cell(row=r, column=2, value=f"=Partite!$B{r}")
            pr.cell(row=r, column=3, value=f'=IF(Partite!$E{r}="","",Partite!$E{r})')
            pr.cell(row=r, column=4, value=f'=IF(Partite!$F{r}="","",Partite!$F{r})')
            for col in range(1, 5):
                s(pr.cell(row=r, column=col), b=(col in (3, 4)), sz=10, fg=MUTE if col == 1 else INK,
                  al="center" if col in (1, 2) else "left", bg=bg, box=BOX)
            for i in range(NP):
                c0 = 5 + i*3
                vals = guess.get(PLAYERS[i]) or [None, None, None]
                for j in range(3):
                    cc = inp(pr.cell(row=r, column=c0+j))
                    if vals[j] is not None: cc.value = vals[j]
                dv.add(pr.cell(row=r, column=c0))
            r += 1
    pr.freeze_panes = "E3"; pr.auto_filter.ref = f"A2:D{RN}"

    # =========================================================== PUNTI
    pn = wb.create_sheet("Punti")
    pn.sheet_view.showGridLines = False; pn.sheet_properties.tabColor = MUTE
    widths(pn, {"A": 9, "B": 8, "C": 28})
    band(pn, "A1:H1", "A1", "H1", bg=NAVY, fg=WHITE, txt="  PUNTI PER PARTITA - tutto calcolato, non scrivere qui", sz=11, h=22)
    s(pn["J1"], sz=8, fg=MUTE, it=True)
    pn["J1"] = ("servizio: a sinistra il segno effettivo (dichiarato, o dedotto dal punteggio), "
                "a destra il punteggio scritto. Servono a vedere chi era da solo.")
    for col, h in zip("ABC", ["ID", "G.", "Partita"]):
        s(pn[f"{col}2"], b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX); pn[f"{col}2"] = h
    # colonne di servizio: i segni effettivi in un blocco, i punteggi scritti in
    # un altro. Devono stare attaccate, perche' le formule ci contano dentro con
    # un COUNTIF per sapere se qualcun altro aveva scelto la stessa cosa.
    HS = 10             # prima colonna dei segni
    HP = HS + NP + 1    # prima colonna dei punteggi
    SEGNI = f"${CL(HS)}{{r}}:${CL(HS+NP-1)}{{r}}"
    PUNTEGGI = f"${CL(HP)}{{r}}:${CL(HP+NP-1)}{{r}}"
    for i in range(NP):
        s(pn.cell(row=2, column=4+i, value=f"=Config!$A${6+i}"), b=True, sz=9, fg=NAVY, bg=PCOL[i], al="center", box=BOX)
        pn.column_dimensions[CL(4+i)].width = 12
        for c0 in (HS, HP):
            s(pn.cell(row=2, column=c0+i, value=f"=Config!$A${6+i}"), b=True, sz=8, fg=MUTE, bg=PANEL2, al="center")
            pn.column_dimensions[CL(c0+i)].width = 9

    r = R0
    for k, g in enumerate(GIORNATE):
        bg = WHITE if k % 2 == 0 else PANEL
        for m in range(1, MPG+1):
            pn.cell(row=r, column=1, value=f"=Partite!$A{r}")
            pn.cell(row=r, column=2, value=f"=Partite!$B{r}")
            pn.cell(row=r, column=3, value=f'=IF(Partite!$E{r}="","",Partite!$E{r}&"  -  "&Partite!$F{r})')
            segni, punteggi = SEGNI.format(r=r), PUNTEGGI.format(r=r)
            # il solitario vale solo dalla giornata scritta in Config
            da_solo = f'$B{r}>=Config!$D$17'
            for i in range(NP):
                sc, hc, ac = CL(5+i*3), CL(6+i*3), CL(7+i*3)
                hp, pp = CL(HS+i), CL(HP+i)
                pn.cell(row=r, column=HS+i, value=(
                    f'=IF(Pronostici!${sc}{r}<>"",UPPER(Pronostici!${sc}{r}),'
                    f'IF(OR(Pronostici!${hc}{r}="",Pronostici!${ac}{r}=""),"",'
                    f'IF(Pronostici!${hc}{r}>Pronostici!${ac}{r},"1",'
                    f'IF(Pronostici!${hc}{r}=Pronostici!${ac}{r},"X","2"))))'))
                pn.cell(row=r, column=HP+i, value=(
                    f'=IF(OR(Pronostici!${hc}{r}="",Pronostici!${ac}{r}=""),"",'
                    # " a " e non "-": con "2-1" il COUNTIF di Excel potrebbe
                    # leggere una data e non trovare piu' nessuno
                    f'Pronostici!${hc}{r}&" a "&Pronostici!${ac}{r})'))
                pn.cell(row=r, column=4+i, value=(
                    f'=IF(Partite!$K{r}<>"Giocata","",'
                    f'IF(AND(Pronostici!${hc}{r}<>"",Pronostici!${ac}{r}<>"",'
                    f'Pronostici!${hc}{r}=Partite!$G{r},Pronostici!${ac}{r}=Partite!$H{r}),'
                    f'IF(AND({da_solo},COUNTIF({punteggi},${pp}{r})=1),Config!$D$9,Config!$D$6),'
                    f'IF(AND(${hp}{r}<>"",${hp}{r}=Partite!$I{r}),'
                    f'IF(AND({da_solo},COUNTIF({segni},${hp}{r})=1),Config!$D$10,Config!$D$7),'
                    f'Config!$D$8)))'))
            for col in range(1, 9):
                s(pn.cell(row=r, column=col), b=(col >= 4), sz=10, fg=MUTE if col == 1 else INK,
                  al="left" if col == 3 else "center", bg=bg, box=BOX)
            for i in range(NP):
                for c0 in (HS, HP):
                    s(pn.cell(row=r, column=c0+i), sz=8, fg=MUTE, al="center", bg=PANEL2)
            r += 1
    pn.freeze_panes = "D3"
    for i in range(NP):
        col = CL(4+i)
        pn.conditional_formatting.add(f"{col}{R0}:{col}{RN}",
            CellIsRule(operator="greaterThanOrEqual", formula=["3"], fill=PatternFill("solid", fgColor=GREEN)))
        pn.conditional_formatting.add(f"{col}{R0}:{col}{RN}",
            CellIsRule(operator="between", formula=["1", "2"], fill=PatternFill("solid", fgColor=AMBER)))

    # =========================================================== CONFIG
    cf = wb.create_sheet("Config")
    cf.sheet_view.showGridLines = False; cf.sheet_properties.tabColor = PANEL2
    widths(cf, {"A": 34, "B": 10, "C": 38, "D": 10, "E": 10, "F": 34})
    band(cf, "A1:F2", "A1", "F2", bg=NAVY, fg=WHITE, txt="  TROFEO OVALMO - REGOLE E CONFIGURAZIONE", sz=15, h=30)
    cf.row_dimensions[2].height = 6

    s(cf["A4"], b=True, sz=12, fg=NAVY); cf["A4"] = "GIOCATORI"
    s(cf["A5"], sz=8, fg=MUTE, it=True); cf["A5"] = "Cambia un nome qui: si aggiorna in tutti i fogli."
    for i, p in enumerate(PLAYERS):
        c = inp(cf.cell(row=6+i, column=1, value=p), sz=11)
        c.fill = PatternFill("solid", fgColor=PCOL[i])

    s(cf["C4"], b=True, sz=12, fg=NAVY); cf["C4"] = "PUNTEGGIO"
    s(cf["C5"], b=True, sz=9, fg=WHITE, bg=NAVY2, box=BOX); cf["C5"] = "Caso"
    s(cf["D5"], b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX); cf["D5"] = "Punti"
    for i, (lab, val) in enumerate([("Risultato esatto (segno + gol)", 3),
                                    ("Solo segno 1X2 corretto", 1),
                                    ("Sbagliato o pronostico mancante", 0),
                                    ("Risultato esatto, e nessun altro lo aveva scritto", SOL_ESATTO),
                                    ("Segno giusto, e nessun altro lo aveva scelto", SOL_SEGNO)]):
        s(cf.cell(row=6+i, column=3, value=lab), sz=10, bg=WHITE if i % 2 == 0 else PANEL, box=BOX)
        inp(cf.cell(row=6+i, column=4, value=val), sz=11)
    s(cf["C11"], sz=8, fg=MUTE, it=True)
    cf["C11"] = "Il risultato esatto non si somma al segno. Gli ultimi due sono il solitario: valgono il doppio."

    s(cf["C13"], b=True, sz=12, fg=NAVY); cf["C13"] = "PARAMETRI"
    for i, (lab, val) in enumerate([("Stagione", "Serie A 2026/27"),
                                    ("Prima giornata in gioco", GIORNATE[0]),
                                    ("Partite per giornata", MPG),
                                    ("Il solitario vale dalla giornata", SOLITARIO_DA)]):
        s(cf.cell(row=14+i, column=3, value=lab), sz=10, bg=WHITE if i % 2 == 0 else PANEL, box=BOX)
        inp(cf.cell(row=14+i, column=4, value=val))

    s(cf["A20"], b=True, sz=12, fg=NAVY); cf["A20"] = "REGOLE DI GIOCO"
    for i, t in enumerate([
     "1.  Pronostici entro il primo fischio della giornata: dopo non valgono piu'.",
     "2.  Chi non manda i pronostici prende 0 su quella giornata, che conta comunque.",
     "3.  Si puo' dare solo il segno 1/X/2: vale 1 punto se corretto, senza bonus risultato.",
     f"4.  Dalla giornata {SOLITARIO_DA} chi indovina da solo vale doppio: {SOL_SEGNO} il segno, {SOL_ESATTO} il risultato esatto.",
     "5.  Parita' in classifica generale: passa avanti chi ha piu' risultati esatti.",
     "6.  Re della giornata = chi fa piu' punti in quella giornata (solo giornate complete).",
    ]):
        c = cf.cell(row=21+i, column=1, value=t); s(c, sz=10, bg=WHITE if i % 2 == 0 else PANEL)
        cf.merge_cells(start_row=21+i, start_column=1, end_row=21+i, end_column=4)
        cf.row_dimensions[21+i].height = 17

    s(cf["A28"], b=True, sz=12, fg=NAVY); cf["A28"] = "COSA SI COMPILA A MANO"
    for i, (a, b) in enumerate([
     ("Le celle gialle", "sono le uniche da toccare. Tutto il resto e' formula."),
     ("Foglio Partite", "Gol Casa e Gol Ospite: il risultato vero della partita."),
     ("Foglio Pronostici", "per ogni giocatore: Segno, Gol C, Gol O."),
     ("Segno o punteggio?", "basta uno dei due: se metti i gol, il segno lo deduce da solo."),
     ("Classifica e Punti", "si aggiornano da soli, non scriverci dentro."),
    ]):
        s(cf.cell(row=29+i, column=1, value=a), b=True, sz=10, bg=WHITE if i % 2 == 0 else PANEL, box=BOX)
        s(cf.cell(row=29+i, column=3, value=b), sz=10, bg=WHITE if i % 2 == 0 else PANEL, box=BOX)
        cf.merge_cells(start_row=29+i, start_column=3, end_row=29+i, end_column=4)

    s(cf["A36"], b=True, sz=12, fg=NAVY); cf["A36"] = "ESEMPIO - Milan-Venezia finisce 2-1"
    for j, h in enumerate(["Giocatore", "Segno", "Gol C", "Gol O", "Punti", "Perche'"]):
        s(cf.cell(row=37, column=1+j, value=h), b=True, sz=9, fg=WHITE, bg=NAVY2, al="center", box=BOX)
    for i, row in enumerate([("Berta", "1", 2, 1, SOL_ESATTO, "risultato esatto, e nessun altro aveva scritto 2-1"),
                             ("Super Gulp", "1", 3, 0, 1, "segno giusto, ma l'1 lo avevano in tre"),
                             ("Lenzuolo", "1", None, None, 1, "solo il segno, giusto ma in compagnia"),
                             ("Just Lele", "X", 1, 1, 0, "segno sbagliato"),
                             ("Lippi", None, None, None, 0, "non ha mandato nulla")]):
        for j, v in enumerate(row):
            c = cf.cell(row=38+i, column=1+j, value=v)
            if 1 <= j <= 3: inp(c, sz=9)
            elif j == 4: s(c, b=True, sz=10, fg=NAVY, al="center", box=BOX,
                           bg=GREEN if v >= 3 else (AMBER if v > 0 else WHITE))
            elif j == 5: s(c, sz=9, fg=MUTE, it=True, box=BOX)
            else: s(c, b=True, sz=9, bg=PCOL[i], box=BOX)
    s(cf["A43"], sz=8, fg=MUTE, it=True)
    cf["A43"] = (f"Esempio dalla giornata {SOLITARIO_DA} in poi. Se il segno giusto lo avesse scelto una "
                 f"persona sola, a lei varrebbe {SOL_SEGNO} punti invece di 1.")

    s(cf["A45"], sz=8, fg=MUTE, it=True)
    cf["A45"] = ("Il file si allunga da solo: quando esce il calendario di una nuova giornata "
                 "viene aggiunto in Partite e le righe di Pronostici e Punti compaiono con lui.")
    s(cf["A46"], sz=8, fg=MUTE, it=True)
    cf["A46"] = "Calendario e risultati: football-data.org, riletti a ogni aggiornamento automatico."

    wb.active = 0
    wb.save(OUT)
    print(f"scritto {OUT} | giornate {GIORNATE[0]}-{GIORNATE[-1]} ({len(GIORNATE)}) | righe {R0}-{RN}")


def verifica(percorso):
    """Riapre il file e controlla che sia sano. Alza AssertionError se non lo e'.

    Un vero ricalcolo richiederebbe Excel o LibreOffice, che sul server di
    GitHub non ci sono. Si controlla quindi cio' che si puo' controllare senza
    aprirlo davvero, ed e' comunque cio' che si rompe in pratica: formule con
    riferimenti morti, fogli spariti, grafico senza dati.
    """
    from openpyxl import load_workbook

    wb = load_workbook(percorso)
    attesi = ["Classifica", "Partite", "Pronostici", "Punti", "Config"]
    mancanti = [f for f in attesi if f not in wb.sheetnames]
    assert not mancanti, f"nel file Excel mancano i fogli: {mancanti}"

    rotte = []
    formule = 0
    for ws in wb.worksheets:
        for riga in ws.iter_rows():
            for cella in riga:
                v = cella.value
                if isinstance(v, str) and v.startswith("="):
                    formule += 1
                    if "#REF!" in v or "#NAME?" in v:
                        rotte.append(f"{ws.title}!{cella.coordinate}: {v[:60]}")
    assert not rotte, "formule rotte nel file Excel:\n  " + "\n  ".join(rotte[:10])
    assert formule > 0, "il file Excel non contiene nessuna formula: qualcosa e' andato storto"

    classifica = wb["Classifica"]
    assert classifica._charts, "manca il grafico dell'andamento nel foglio Classifica"
    return dict(formule=formule, fogli=wb.sheetnames)


if __name__ == "__main__":
    # uso a mano:  python -m ovalmo.excel dati.json Trofeo_Ovalmo.xlsx
    dati = json.load(open(sys.argv[1], encoding="utf-8"))
    uscita = sys.argv[2] if len(sys.argv) > 2 else "Trofeo_Ovalmo.xlsx"
    genera(dati, uscita)
    print(verifica(uscita))
